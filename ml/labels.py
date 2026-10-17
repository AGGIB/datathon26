"""Loader for the organizers' open_labels.xlsx (video name, timestamp, violation flag)."""

import re
from dataclasses import dataclass
from pathlib import Path

_NUM = r"(\d+(?:[.,]\d+)?)"
_MIN_RE = re.compile(_NUM + r"\s*мин\w*")
_SEC_RE = re.compile(_NUM + r"\s*сек\w*")
_CLOCK_RE = re.compile(r"^\s*(\d{1,2}):(\d{2}(?:[.,]\d+)?)\s*$")
_CYRILLIC_X = str.maketrans({"Х": "X", "х": "X", "x": "X"})


@dataclass(frozen=True)
class Label:
    video: str
    has_violation: bool
    t_sec: float | None


def _num(text: str) -> float:
    return float(text.replace(",", "."))


def parse_timestamp(raw: str | float | int | None) -> float | None:
    """'22 секунда' -> 22.0, '1минута 8 секунда' -> 68.0, '8,5 секунды' -> 8.5, '01:08' -> 68.0, 57 -> 57.0."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        if raw < 0:
            raise ValueError(f"negative timestamp: {raw!r}")
        return float(raw)
    text = str(raw).lower()
    clock = _CLOCK_RE.match(text)
    if clock:
        minutes, seconds = int(clock.group(1)), _num(clock.group(2))
    else:
        if "-" in text:
            raise ValueError(f"negative timestamp: {raw!r}")
        m, s = _MIN_RE.search(text), _SEC_RE.search(text)
        if not m and not s:
            raise ValueError(f"unrecognized timestamp: {raw!r}")
        minutes, seconds = (_num(m.group(1)) if m else 0), (_num(s.group(1)) if s else 0.0)
    if minutes and seconds >= 60:
        raise ValueError(f"seconds out of range in {raw!r}")
    return float(minutes * 60 + seconds)


def normalize_video_name(raw: str) -> str:
    """Dataset mixes Latin X and Cyrillic Х; we key everything by 'X-<n>'."""
    stem = Path(raw.strip()).stem
    return stem.translate(_CYRILLIC_X).upper()


def load_labels(xlsx_path: Path) -> list[Label]:
    from openpyxl import load_workbook

    sheet = load_workbook(xlsx_path, read_only=True).active
    labels = []
    for name, stamp, flag in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        if not name:
            continue
        has_violation = str(flag).strip().lower() == "нарушение"
        labels.append(Label(normalize_video_name(str(name)), has_violation, parse_timestamp(stamp)))
    return labels
