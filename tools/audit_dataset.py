#!/usr/bin/env python3
"""Read-only intake: labels, hashes, taxonomy and structural IDs, never video content.

Requires openpyxl and python-docx. No network, frame decoding or package-script imports.
Only the JSON report is written, strictly under this repository's data/audit/.
"""

import argparse
import ast
from collections import Counter, defaultdict
from datetime import datetime, time, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import unicodedata


ROOT = Path(__file__).resolve().parents[1]
HEADERS = ("наименование файла", "таймстемп", "нарушение / без нарушения")
LABELS = {"нарушение": "violation", "без нарушения": "no_violation"}
VIDEO_NAME = re.compile(r"[XХxх]-\d+(?:\.mov)?", re.IGNORECASE)
TYPE_HEADING = re.compile(r"([ТД]\d{2})\s+(.+)")
REPORT_WIDTHS = {"sources": 4, "types": 8, "additional_checks": 6}
REGISTRY_FIELDS = {"id", "name", "original_label", "timestamp", "summary", "assessment", "observation"}
VISUAL_FIELDS = {"id", "name", "evidence", "hypothesis", "limitations", "observation", "observed_interval", "visual_confidence"}


def nfc(value):
    return unicodedata.normalize("NFC", value)


def whitespace(value):
    return " ".join(nfc(value).split())


def name_key(value):
    """Comparison key only. Each namespace must reject collisions before joining."""
    if not isinstance(value, str) or not VIDEO_NAME.fullmatch(nfc(value)):
        raise ValueError("invalid_technical_video_name")
    stem = re.sub(r"\.mov$", "", nfc(value), flags=re.IGNORECASE)
    return stem.casefold().replace("х", "x")


def parse_timestamp(value):
    """Seconds, mm:ss, or Russian minute/second units. No fuzzy guessing."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("invalid_timestamp_type")
    if isinstance(value, time):
        result = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    elif isinstance(value, timedelta):
        result = value.total_seconds()
    elif isinstance(value, (int, float)):
        result = float(value)
    elif isinstance(value, str):
        text = whitespace(value).casefold().replace(",", ".")
        number = r"\d+(?:\.\d+)?"
        minutes = r"(?:минута|минуты|минуту|минут|мин\.?|м\.?)"
        seconds = r"(?:секунда|секунды|секунду|секунд|сек\.?|с\.?)"
        if re.fullmatch(number, text):
            result = float(text)
        elif match := re.fullmatch(r"(\d+):(\d{2}(?:\.\d+)?)", text):
            if float(match[2]) >= 60:
                raise ValueError("invalid_timestamp_seconds_component")
            result = int(match[1]) * 60 + float(match[2])
        elif match := re.fullmatch(rf"({number})\s*{minutes}(?:\s*({number})\s*{seconds})?", text):
            if match[2] is not None and float(match[2]) >= 60:
                raise ValueError("invalid_timestamp_seconds_component")
            result = float(match[1]) * 60 + (float(match[2]) if match[2] is not None else 0)
        elif match := re.fullmatch(rf"({number})\s*{seconds}", text):
            result = float(match[1])
        else:
            raise ValueError("unrecognized_timestamp")
    else:
        raise ValueError("invalid_timestamp_type")
    if not math.isfinite(result) or result < 0:
        raise ValueError("invalid_timestamp_range")
    return result


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def json_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("duplicate_json_key")
        result[key] = value
    return result


def read_json(path):
    def reject_constant(_):
        raise ValueError("nonfinite_json_number")
    return json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=json_object,
                      parse_constant=reject_constant)


def input_path(value):
    path = Path(value)
    path = path if path.is_absolute() else ROOT / path
    if path.exists():
        return path.resolve()
    # macOS filenames may be stored in NFD; never choose between NFC collisions.
    matches = [p for p in path.parent.iterdir() if nfc(p.name) == nfc(path.name)] if path.parent.is_dir() else []
    if len(matches) != 1:
        raise ValueError("input_missing_or_nfc_ambiguous")
    return matches[0].resolve()


def path_label(path):
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)


class Audit:
    def __init__(self):
        self.issues = []
        self.name_checks = {}

    def issue(self, code, where, severity="error", **details):
        self.issues.append(dict(code=code, where=where, severity=severity, **details))

    def index_names(self, rows, where):
        groups = defaultdict(list)
        for row in rows:
            try:
                key = name_key(row["name"])
            except (KeyError, ValueError):
                self.issue("invalid_video_name", where, position=row.get("position"))
                continue
            groups[key].append(row)
        collisions = []
        for key, group in groups.items():
            if len(group) > 1:
                raw = [r["name"] for r in group]
                kind = "exact_duplicate" if len(set(raw)) == 1 else (
                    "nfc_collision" if len({nfc(s) for s in raw}) == 1 else "normalized_name_collision")
                collisions.append(key)
                self.issue(kind, where, key=key, names=raw)
        self.name_checks[where] = {"rows": len(rows), "unique_keys": len(groups),
                                  "collision_groups": len(collisions)}
        # Ambiguous groups are deliberately absent from the join index.
        return {key: group[0] for key, group in groups.items() if len(group) == 1}

    def compare_keys(self, left, right, where):
        only_left = sorted(set(left) - set(right))
        only_right = sorted(set(right) - set(left))
        if only_left or only_right:
            self.issue("id_set_mismatch", where, only_left=only_left, only_right=only_right)
        return {"matched": len(set(left) & set(right)), "only_left": only_left, "only_right": only_right}

    def count(self, actual, expected, where):
        if actual != expected:
            self.issue("unexpected_count", where, actual=actual, expected=expected)


def read_labels(path, sheet, audit):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows or tuple(whitespace(str(x)).casefold() for x in rows[0]) != HEADERS:
        raise ValueError("invalid_labels_header")
    result = []
    for position, values in enumerate(rows[1:], start=2):
        if all(x is None for x in values):
            audit.issue("blank_label_row", "labels", position=position)
            continue
        if len(values) != 3:
            audit.issue("invalid_label_row_width", "labels", position=position)
            continue
        label = LABELS.get(whitespace(str(values[2])).casefold())
        if label is None:
            audit.issue("unknown_label", "labels", position=position)
        try:
            seconds = parse_timestamp(values[1])
        except ValueError as error:
            audit.issue(str(error), "labels", position=position)
            seconds = None
        if label == "violation" and seconds is None:
            audit.issue("missing_violation_timestamp", "labels", position=position)
        if label == "no_violation" and seconds is not None:
            audit.issue("timestamp_for_negative_label", "labels", position=position)
        # Never store any narrative; these are the three authorized XLSX fields.
        result.append({"position": position, "name": values[0], "label": label,
                       "timestamp_original": str(values[1]) if values[1] is not None else None,
                       "timestamp_seconds": seconds})
    return result


def manifest_rows(raw, audit):
    if not isinstance(raw, list):
        raise ValueError("manifest_must_be_array")
    result = []
    ids = []
    for position, item in enumerate(raw, start=1):
        if not isinstance(item, dict) or not isinstance(item.get("file"), str):
            audit.issue("invalid_manifest_row", "manifest", position=position)
            continue
        # The author's absolute path is not followed. Only its technical basename is used.
        filename = item["file"].replace("\\", "/").rsplit("/", 1)[-1]
        row = {"position": position, "name": filename, "id": item.get("id"),
               "sha256": item.get("sha256"), "duration": item.get("duration"),
               "label": item.get("label"), "timestamp": item.get("timestamp")}
        for key in ("size_bytes", "size", "bytes"):
            if key in item:
                row["size_bytes"] = item[key]
                break
        try:
            key = name_key(filename)
            if name_key(item.get("name")) != key:
                audit.issue("manifest_name_mismatch", "manifest", position=position)
            if type(row["id"]) is not int or row["id"] != int(key.split("-")[1]):
                audit.issue("manifest_id_mismatch", "manifest", position=position)
        except ValueError:
            audit.issue("invalid_manifest_name", "manifest", position=position)
        if type(row["id"]) is int:
            ids.append(row["id"])
        result.append(row)
    if len(ids) != len(set(ids)):
        audit.issue("duplicate_numeric_id", "manifest")
    return result


def audit_video_files(videos, labels, manifest, audit):
    records = []
    for key, video in sorted(videos.items()):
        path = video["path"]
        before = path.stat()
        digest = sha256(path)
        after = path.stat()
        stable = (before.st_size, before.st_mtime_ns) == (after.st_size, after.st_mtime_ns)
        if not stable:
            audit.issue("file_changed_during_hash", "videos", key=key)
        source = manifest.get(key)
        label = labels.get(key)
        hash_status, size_status = "not_provided", "not_provided"
        duration, duration_status = None, "unverified_no_manifest"
        if source:
            expected = source.get("sha256")
            if expected is None:
                audit.issue("missing_manifest_hash", "manifest", severity="warning", key=key)
            elif not isinstance(expected, str) or not re.fullmatch(r"[0-9a-fA-F]{64}", expected):
                hash_status = "invalid"
                audit.issue("invalid_manifest_hash", "manifest", key=key)
            else:
                hash_status = "matched" if stable and expected.lower() == digest else "mismatch"
                if hash_status == "mismatch":
                    audit.issue("hash_mismatch", "videos", key=key)
            if "size_bytes" in source:
                size = source["size_bytes"]
                size_status = "matched" if type(size) is int and size == after.st_size else "mismatch"
                if size_status == "mismatch":
                    audit.issue("size_mismatch", "videos", key=key)
            value = source.get("duration")
            if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) and value > 0:
                duration = float(value)
                duration_status = "manifest_hash_matched" if hash_status == "matched" else "unverified_hash_" + hash_status
            else:
                duration_status = "unverified_invalid_duration"
                audit.issue("invalid_manifest_duration", "manifest", key=key)
            if label:
                source_label = LABELS.get(whitespace(str(source.get("label"))).casefold())
                if source_label != label["label"]:
                    audit.issue("manifest_label_mismatch", "manifest", key=key)
                try:
                    source_time = parse_timestamp(source.get("timestamp"))
                except ValueError:
                    source_time = "invalid"
                if source_time != label["timestamp_seconds"]:
                    audit.issue("manifest_timestamp_mismatch", "manifest", key=key)
        bounds = "not_applicable" if label and label["timestamp_seconds"] is None else "unverified"
        if label and label["timestamp_seconds"] is not None and duration_status == "manifest_hash_matched":
            bounds = "within_manifest_duration" if 0 <= label["timestamp_seconds"] < duration else "out_of_bounds"
            if bounds == "out_of_bounds":
                audit.issue("timestamp_out_of_bounds", "labels", key=key)
        records.append({"key": key, "filename": path.name, "size_bytes": after.st_size,
                        "sha256": digest, "manifest_hash_check": hash_status,
                        "manifest_size_check": size_status,
                        "duration_seconds": duration, "duration_provenance": "author_manifest",
                        "duration_verification": duration_status, "duration_measured_by_audit": False,
                        "label": label["label"] if label else None,
                        "timestamp_seconds": label["timestamp_seconds"] if label else None,
                        "timestamp_original": label.get("timestamp_original") if label else None,
                        "timestamp_bounds": bounds,
                        "name_spellings": {"video": video["name"], "labels": label["name"] if label else None,
                                           "manifest": source["name"] if source else None}})
    hashes = defaultdict(list)
    for record in records:
        hashes[record["sha256"]].append(record["key"])
    by_key = {record["key"]: record for record in records}
    for keys in hashes.values():
        if len(keys) > 1:
            audit.issue("duplicate_file_content", "videos", severity="warning", keys=keys)
            annotations = [{"key": key, "label": by_key[key]["label"],
                            "timestamp_seconds": by_key[key]["timestamp_seconds"],
                            "timestamp_original": by_key[key]["timestamp_original"]} for key in keys]
            if len({row["label"] for row in annotations}) > 1:
                audit.issue("duplicate_content_label_conflict", "videos", severity="warning", annotations=annotations)
            if len({row["timestamp_seconds"] for row in annotations}) > 1:
                audit.issue("duplicate_content_multiple_timestamps", "videos", severity="warning", annotations=annotations)
    return records


def read_docx(path, audit):
    from docx import Document
    doc = Document(path)
    types, blocks = [], {}
    current = None
    for paragraph in doc.paragraphs:
        text = paragraph.text
        match = TYPE_HEADING.fullmatch(whitespace(text)) if paragraph.style.name == "Heading 2" else None
        if match:
            current = match[1]
            types.append({"id": current, "name": match[2]})
            if current in blocks:
                audit.issue("duplicate_taxonomy_id", "docx", id=current)
            blocks[current] = []
        elif paragraph.style.name.startswith("Heading"):
            current = None
        elif current:
            # Used locally to compare generic references. Never serialized or printed.
            blocks[current].append(text)
    registry = []
    matching_tables = []
    for table in doc.tables:
        if table.rows and whitespace(table.rows[0].cells[0].text) == "Видео и метка":
            matching_tables.append(table)
    if len(matching_tables) != 1:
        raise ValueError("docx_registry_table_missing_or_ambiguous")
    table = matching_tables[0]
    if len(table.columns) != 4:
        audit.issue("docx_registry_column_count", "docx")
    for position, row in enumerate(table.rows[1:], start=2):
        # No reads of observation/assessment columns or per-video paragraphs.
        match = re.match(r"^([XХxх]-\d+(?:\.mov)?)(?=\s|$)", row.cells[0].text.strip(), re.I)
        if match:
            registry.append({"position": position, "name": match[1]})
        else:
            audit.issue("docx_registry_invalid_id", "docx", position=position)
    return types, blocks, registry


def taxonomy_audit(raw, doc_types, blocks, doc_registry, labels, expected, audit):
    if not isinstance(raw, dict) or set(raw) != {*REPORT_WIDTHS, "registry"}:
        raise ValueError("invalid_report_sections")
    for section, width in REPORT_WIDTHS.items():
        rows = raw[section]
        if not isinstance(rows, list):
            raise ValueError("report_section_must_be_array")
        for row in rows:
            if not isinstance(row, list) or len(row) != width:
                raise ValueError("report_row_array_width_mismatch")
            if not all(isinstance(x, str) for x in row[:width - (section != "sources")]):
                raise ValueError("report_row_field_type_mismatch")
            if section != "sources" and (not isinstance(row[-1], list) or not all(isinstance(x, str) for x in row[-1])):
                raise ValueError("report_source_ids_must_be_array")
    sources = [{"id": r[0], "title": r[1], "url": r[3], "legal_review": "not_reviewed"} for r in raw["sources"]]
    source_ids = [r["id"] for r in sources]
    if len(source_ids) != len(set(source_ids)):
        audit.issue("duplicate_source_id", "report_data")
    doc_index = {r["id"]: r for r in doc_types}
    generic = []
    comparisons = {}
    for section, prefix, count, legal_index in (("types", "Т", 17, 4), ("additional_checks", "Д", 15, 2)):
        rows = raw[section]
        ids = [r[0] for r in rows]
        expected_ids = {f"{prefix}{i:02}" for i in range(1, count + 1)}
        audit.count(len(rows), count, section)
        if len(ids) != len(set(ids)):
            audit.issue("duplicate_taxonomy_id", section)
        comparisons[section] = audit.compare_keys(set(ids), expected_ids, section + "/expected_ids")
        audit.compare_keys(set(ids), {i for i in doc_index if i.startswith(prefix)}, section + "/docx_ids")
        for row in rows:
            tid, name, refs = row[0], row[1], row[-1]
            block = whitespace("\n".join(blocks.get(tid, [])))
            doc_refs = sorted(set(re.findall(r"\bИ\d+\b", block)))
            name_matches = tid in doc_index and whitespace(name) == whitespace(doc_index[tid]["name"])
            reference_matches = whitespace(row[legal_index]) in block
            if not name_matches:
                audit.issue("taxonomy_name_mismatch", section, id=tid)
            if not reference_matches or set(refs) != set(doc_refs):
                audit.issue("taxonomy_reference_mismatch", section, id=tid)
            unknown = sorted(set(refs) - set(source_ids))
            if unknown:
                audit.issue("unknown_source_reference", section, id=tid, source_ids=unknown)
            generic.append({"id": tid, "name": name, "source_ids": refs,
                            "docx_source_ids": doc_refs, "name_matches_docx": name_matches,
                            "legal_reference_matches_docx": reference_matches, "legal_review": "not_reviewed"})
    registry = raw["registry"]
    if not isinstance(registry, list):
        raise ValueError("report_registry_must_be_array")
    technical_registry = []
    numeric_ids = []
    for position, row in enumerate(registry, start=1):
        if not isinstance(row, dict) or set(row) != REGISTRY_FIELDS:
            raise ValueError("invalid_report_registry_schema")
        # Only structural types and technical identifiers are inspected here.
        if not (type(row["id"]) is int and isinstance(row["name"], str)
                and isinstance(row["original_label"], str) and isinstance(row["summary"], str)
                and (row["timestamp"] is None or type(row["timestamp"]) in (int, float))
                and (row["assessment"] is None or isinstance(row["assessment"], list))
                and isinstance(row["observation"], dict)):
            audit.issue("invalid_registry_field_type", "report_registry", position=position)
        technical_registry.append({"position": position, "name": row["name"]})
        numeric_ids.append(row["id"])
        try:
            if row["id"] != int(name_key(row["name"]).split("-")[1]):
                audit.issue("registry_numeric_id_mismatch", "report_registry", position=position)
        except ValueError:
            audit.issue("registry_invalid_name", "report_registry", position=position)
    if len(numeric_ids) != len(set(numeric_ids)):
        audit.issue("duplicate_numeric_id", "report_registry")
    audit.count(len(registry), expected, "report_registry")
    audit.count(len(doc_registry), expected, "docx_registry")
    report_index = audit.index_names(technical_registry, "report_registry")
    doc_registry_index = audit.index_names(doc_registry, "docx_registry")
    registry_comparison = audit.compare_keys(report_index, doc_registry_index, "registry/docx_vs_json")
    audit.compare_keys(report_index, labels, "registry/labels")
    return {"types_count": len(raw["types"]), "additional_checks_count": len(raw["additional_checks"]),
            "sources": sources, "generic_entries": generic, "id_comparisons": comparisons,
            "registry_comparison": registry_comparison, "registry_ids": sorted(report_index),
            "registry_narrative_exported": False}


def visual_ids(paths, labels, expected, audit):
    combined, counts = [], {}
    for path in paths:
        rows = read_json(path)
        if not isinstance(rows, list):
            raise ValueError("visual_file_must_be_array")
        counts[path.name] = len(rows)
        for position, row in enumerate(rows, start=1):
            if not isinstance(row, dict) or set(row) != VISUAL_FIELDS:
                raise ValueError("invalid_visual_schema")
            try:
                if type(row["id"]) is not int or row["id"] != int(name_key(row["name"]).split("-")[1]):
                    audit.issue("visual_numeric_id_mismatch", path.name, position=position)
            except ValueError:
                audit.issue("invalid_visual_name", path.name, position=position)
            combined.append({"position": position, "name": row["name"]})
    audit.count(len(combined), expected, "visual_ids")
    index = audit.index_names(combined, "visual_ids")
    comparison = audit.compare_keys(index, labels, "visual_ids/labels")
    return {"counts": counts, "comparison": comparison, "narrative_exported": False}


def static_package_inventory(folder, audit):
    scripts = []
    for path in sorted(folder.glob("*.py")):
        try:
            tree = ast.parse(path.read_text(encoding="utf-8"))
        except (SyntaxError, UnicodeError):
            audit.issue("package_script_ast_error", path.name, severity="warning")
            continue
        constants = [n.value for n in ast.walk(tree) if isinstance(n, ast.Constant) and isinstance(n.value, str)]
        scripts.append({"file": path.name, "sha256": sha256(path),
                        "absolute_user_path_literals": sum(s.startswith(("/Users/", "/home/")) for s in constants),
                        "parsed_only": True, "executed": False})
    return {"extension_counts": dict(Counter(p.suffix.lower() for p in folder.rglob("*") if p.is_file())),
            "scripts": scripts, "media_opened": False, "scripts_executed": False}


def run(args):
    audit = Audit()
    paths = {name: input_path(getattr(args, name)) for name in
             ("labels", "videos", "manifest", "taxonomy", "report_data", "visual_first", "visual_second")}
    input_hashes = {name: sha256(path) for name, path in paths.items() if path.is_file()}
    labels = read_labels(paths["labels"], args.sheet, audit)
    video_rows = [{"name": p.name, "path": p} for p in sorted(paths["videos"].iterdir())
                  if p.is_file() and p.suffix.casefold() == ".mov"]
    manifest = manifest_rows(read_json(paths["manifest"]), audit)
    for where, rows in (("labels", labels), ("videos", video_rows), ("manifest", manifest)):
        audit.count(len(rows), args.expected_count, where)
    label_index = audit.index_names(labels, "labels")
    video_index = audit.index_names(video_rows, "videos")
    manifest_index = audit.index_names(manifest, "manifest")
    joins = {"labels_videos": audit.compare_keys(label_index, video_index, "labels/videos"),
             "videos_manifest": audit.compare_keys(video_index, manifest_index, "videos/manifest")}
    videos = audit_video_files(video_index, label_index, manifest_index, audit)
    doc_types, blocks, doc_registry = read_docx(paths["taxonomy"], audit)
    taxonomy = taxonomy_audit(read_json(paths["report_data"]), doc_types, blocks, doc_registry,
                              label_index, args.expected_count, audit)
    visual = visual_ids([paths["visual_first"], paths["visual_second"]], label_index, args.expected_count, audit)
    package = static_package_inventory(paths["manifest"].parent, audit)
    for name, original_hash in input_hashes.items():
        if sha256(paths[name]) != original_hash:
            audit.issue("input_changed_during_audit", name)
    summary = {"label_rows": len(labels), "video_files": len(video_rows), "manifest_rows": len(manifest),
               "unique_video_contents": len({r["sha256"] for r in videos}),
               "labels": dict(Counter(r["label"] or "unknown" for r in labels)),
               "hash_checks": dict(Counter(r["manifest_hash_check"] for r in videos)),
               "size_checks": dict(Counter(r["manifest_size_check"] for r in videos)),
               "duration_verification": dict(Counter(r["duration_verification"] for r in videos)),
               "timestamp_bounds": dict(Counter(r["timestamp_bounds"] for r in videos)),
               "video_filename_alphabets": dict(Counter("Cyrillic_Х" if "Х" in nfc(r["name"]) else "Latin_X" for r in video_rows)),
               "types": taxonomy["types_count"], "additional_checks": taxonomy["additional_checks_count"],
               "registry_ids_matched": taxonomy["registry_comparison"]["matched"],
               "generic_names_matched": sum(r["name_matches_docx"] for r in taxonomy["generic_entries"]),
               "generic_legal_refs_matched": sum(r["legal_reference_matches_docx"] for r in taxonomy["generic_entries"]),
               "errors": sum(i["severity"] == "error" for i in audit.issues),
               "warnings": sum(i["severity"] == "warning" for i in audit.issues),
               "issue_codes": dict(Counter(i["code"] for i in audit.issues))}
    return {"schema_version": 1, "audited_at_utc": datetime.now(timezone.utc).isoformat(),
            "inputs": {k: {"path": path_label(v), "sha256": input_hashes.get(k)} for k, v in paths.items()},
            "summary": summary, "name_checks": audit.name_checks, "joins": joins,
            "label_rows": labels, "videos": videos, "taxonomy": taxonomy,
            "visual_files_structure": visual, "package_static": package, "issues": audit.issues,
            "limitations": ["No video, audio or frames decoded; no visual validation.",
                            "Duration is author manifest metadata, never measured by this audit.",
                            "Author taxonomy and legal references are not reviewed ground truth.",
                            "Approximate timestamps were parsed, not visually refined."]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    defaults = {"labels": "data/labels/open/open_labels.xlsx", "videos": "data/labels/open/videos",
                "manifest": "data/videos/analysis/manifest.json", "taxonomy": "data/labels/Типы нарушений.docx",
                "report_data": "data/videos/analysis/report_data.json", "visual_first": "data/videos/analysis/visual_first.json",
                "visual_second": "data/videos/analysis/visual_second.json"}
    for name, default in defaults.items():
        parser.add_argument("--" + name.replace("_", "-"), default=default)
    parser.add_argument("--sheet", default="labels")
    parser.add_argument("--expected-count", type=int, default=42)
    parser.add_argument("--output", default="data/audit/intake.json")
    args = parser.parse_args()
    output = Path(args.output)
    output = (output if output.is_absolute() else ROOT / output).resolve()
    allowed = ROOT / "data" / "audit"
    if not output.is_relative_to(allowed) or output.suffix != ".json" or args.expected_count <= 0:
        parser.error("output must be a .json file under repository data/audit; expected-count must be positive")
    try:
        result = run(args)
        if any(output == input_path(getattr(args, name)) for name in defaults):
            raise ValueError("output_overlaps_input")
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    except (OSError, ValueError, KeyError, ImportError) as error:
        # Do not echo arbitrary input content, JSON excerpts or document text.
        detail = str(error) if re.fullmatch(r"[a-z_]+", str(error)) else type(error).__name__
        print(f"AUDIT FAILED: {detail}; check input paths/schema/dependencies.", file=sys.stderr)
        return 2
    print(json.dumps({"summary": result["summary"], "report": path_label(output)}, ensure_ascii=False, indent=2))
    return 1 if result["summary"]["errors"] else 0


if __name__ == "__main__":
    sys.exit(main())
