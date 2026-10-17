"""Synthetic audit inputs only: byte fixtures are deliberately not real videos."""

import copy
import hashlib
import json

from docx import Document
from openpyxl import Workbook, load_workbook
import pytest

from ml import labels as shared
from tools import audit_dataset as audit


def save_labels(path, rows):
    book = Workbook()
    book.active.title = "labels"
    book.active.append(list(audit.HEADERS))
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()


@pytest.fixture
def dataset(tmp_path, monkeypatch):
    monkeypatch.setattr(audit, "ROOT", tmp_path)
    folder = tmp_path / "data/open"
    videos = folder / "videos"
    videos.mkdir(parents=True)
    (videos / "X-1.mov").write_bytes(b"synthetic non-video fixture one")
    (videos / "Х-2.MOV").write_bytes(b"synthetic non-video fixture two")
    labels = folder / "open_labels.xlsx"
    save_labels(
        labels, [["Х-1", "8.5 секунд", "нарушение"], ["X-2", None, "без нарушения"]]
    )
    return tmp_path, labels, videos


def invoke(dataset, *args):
    root, _, _ = dataset
    code = audit.main(["--expected-count", "2", *map(str, args)])
    report = root / "data/audit/intake.json"
    return code, json.loads(
        report.read_text()
    ) if report.exists() and not report.is_symlink() else None


@pytest.fixture
def author_package(dataset):
    root, _, videos = dataset
    folder = root / "author"
    folder.mkdir()
    source = ["И1", "Synthetic act", "synthetic", "https://example.test/act"]
    raw = {"sources": [source], "types": [], "additional_checks": [], "registry": []}
    doc = Document()
    doc.add_heading("Synthetic general taxonomy", 1)
    for section, prefix, count in [("types", "Т", 17), ("additional_checks", "Д", 15)]:
        for number in range(1, count + 1):
            code = f"{prefix}{number:02}"
            name = f"Synthetic {code}"
            doc.add_heading(f"{code} {name}", 2)
            doc.add_paragraph("Synthetic law И1")
            raw[section].append(
                [
                    code,
                    name,
                    "cue",
                    "scope",
                    "Synthetic law И1",
                    "context",
                    "limit",
                    ["И1"],
                ]
                if section == "types"
                else [code, name, "Synthetic law И1", "scope", "limit", ["И1"]]
            )
    doc.add_heading("Registry", 1)
    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Видео и метка"
    manifest, visual = [], []
    for number, path in enumerate(sorted(videos.iterdir()), 1):
        name = f"X-{number}"
        flag, stamp = ("нарушение", 8.5) if number == 1 else ("без нарушения", None)
        table.add_row().cells[0].text = name + " synthetic"
        manifest.append(
            {
                "id": number,
                "file": "/never-follow-this-author-path/" + path.name,
                "name": name,
                "label": flag,
                "timestamp": stamp,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "duration": 12.0,
                "size_bytes": path.stat().st_size,
            }
        )
        raw["registry"].append(
            {
                "id": number,
                "name": name,
                "original_label": flag,
                "timestamp": stamp,
                "summary": "SYNTHETIC_NARRATIVE_NOT_EXPORTED",
                "assessment": [],
                "observation": {},
            }
        )
        visual.append(
            {
                "id": number,
                "name": name,
                "evidence": [],
                "hypothesis": "synthetic",
                "limitations": [],
                "observation": {},
                "observed_interval": None,
                "visual_confidence": 0.0,
            }
        )
    doc.save(folder / "taxonomy.docx")
    (folder / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    (folder / "report.json").write_text(json.dumps(raw), encoding="utf-8")
    for number in range(2):
        (folder / f"visual{number}.json").write_text(
            json.dumps(visual[number : number + 1]), encoding="utf-8"
        )
    (folder / "foreign.py").write_text(
        "raise RuntimeError('MUST_NOT_EXECUTE')\npath='/Users/synthetic'\n"
    )
    return folder, raw, manifest


def package_args(folder):
    return [
        "--manifest",
        folder / "manifest.json",
        "--taxonomy",
        folder / "taxonomy.docx",
        "--report-data",
        folder / "report.json",
        "--visual-first",
        folder / "visual0.json",
        "--visual-second",
        folder / "visual1.json",
    ]


def test_default_readme_core_and_shared_loader(dataset, capsys):
    assert audit.parse_timestamp is shared.parse_timestamp
    assert audit.normalize_video_name is shared.normalize_video_name
    assert audit.load_labels is shared.load_labels
    code, report = invoke(dataset)
    assert code == 0
    assert report["label_rows"][0]["timestamp_seconds"] == 8.5
    assert all(s["status"] == "skipped" for s in report["sections"].values())
    output = capsys.readouterr().out
    assert "X-1" not in output and "Х-2" not in output
    assert report["summary"]["unique_video_contents"] == 2


def test_full_synthetic_xlsx_docx_exit_zero(dataset, author_package):
    folder, _, _ = author_package
    code, report = invoke(dataset, *package_args(folder))
    assert code == 0
    assert report["summary"]["types"] == 17
    assert report["summary"]["additional_checks"] == 15
    assert report["summary"]["generic_names_matched"] == 32
    assert report["summary"]["generic_legal_refs_matched"] == 32
    assert report["taxonomy"]["registry_narrative_exported"] is False
    assert "SYNTHETIC_NARRATIVE_NOT_EXPORTED" not in json.dumps(report)
    assert report["package_static"]["scripts_executed"] is False
    assert report["summary"]["timestamp_bounds"]["within_manifest_duration"] == 1


@pytest.mark.parametrize("conflict", [False, True])
def test_duplicate_exit_and_warning_aggregates(dataset, conflict, capsys):
    _, xlsx, videos = dataset
    (videos / "Х-2.MOV").write_bytes((videos / "X-1.mov").read_bytes())
    save_labels(
        xlsx,
        [
            ["X-1", 8, "нарушение"],
            [
                "Х-2",
                None if conflict else 9,
                "без нарушения" if conflict else "нарушение",
            ],
        ],
    )
    code, report = invoke(dataset)
    assert code == int(conflict)
    assert report["summary"]["issue_codes"]["duplicate_file_content"] == 1
    assert (
        report["summary"]["issue_codes"]["duplicate_content_multiple_timestamps"] == 1
    )
    if conflict:
        assert (
            next(
                i
                for i in report["issues"]
                if i["code"] == "duplicate_content_label_conflict"
            )["severity"]
            == "error"
        )
    output = capsys.readouterr().out
    assert "X-1" not in output and "Х-2" not in output


def test_blank_rows_warn_partial_rows_error(dataset):
    _, xlsx, _ = dataset
    book = load_workbook(xlsx)
    book.active.append([None, None, None])
    book.active.append([" ", " ", " "])
    book.save(xlsx)
    code, report = invoke(dataset)
    assert code == 0 and report["summary"]["issue_codes"]["blank_label_row"] == 2
    book.active.append([None, 3, "нарушение"])
    book.save(xlsx)
    code, report = invoke(dataset)
    assert code == 1 and report["summary"]["issue_codes"]["partial_label_row"] == 1


def test_extra_extensions_are_counted_without_names(dataset, capsys):
    _, _, videos = dataset
    (videos / "private_filename.mp4").write_bytes(b"fixture")
    (videos / "private_notes.txt").write_text("synthetic")
    code, report = invoke(dataset)
    assert code == 0
    assert report["summary"]["unexpected_video_extensions"] == {".mp4": 1, ".txt": 1}
    assert "private_filename" not in capsys.readouterr().out
    assert "private_filename" not in json.dumps(report["issues"])


@pytest.mark.parametrize(
    "value", ["../escape.json", "data/audit/../escape.json", "data/audit/out.txt"]
)
def test_output_traversal_and_suffix_exit_two(dataset, value):
    assert invoke(dataset, "--output", value)[0] == 2
    assert not (dataset[0] / "escape.json").exists()


def test_output_absolute_and_symlink_guards(dataset):
    root = dataset[0]
    outside = root / "outside"
    outside.mkdir()
    assert invoke(dataset, "--output", outside / "escape.json")[0] == 2
    (root / "data/audit").symlink_to(outside, target_is_directory=True)
    assert invoke(dataset)[0] == 2
    assert list(outside.iterdir()) == []


def test_output_nested_success_and_final_symlink(dataset):
    root = dataset[0]
    assert invoke(dataset, "--output", "data/audit/sub/dir/report.json")[0] == 0
    target = root / "target.json"
    target.write_text("unchanged")
    (root / "data/audit/intake.json").symlink_to(target)
    assert invoke(dataset)[0] == 2
    assert target.read_text() == "unchanged"


def test_output_cannot_alias_or_be_inside_input(dataset):
    root, _, videos = dataset
    out = root / "data/audit"
    out.mkdir()
    manifest = out / "manifest.json"
    manifest.write_text("[]")
    assert invoke(dataset, "--manifest", manifest, "--output", manifest)[0] == 2
    with pytest.raises(audit.AuditFailure):
        audit.validate_output(out / "new.json", [out])
    assert manifest.read_text() == "[]"
    assert all(p.suffix.lower() == ".mov" for p in videos.iterdir())


@pytest.mark.parametrize(
    "args",
    [
        ["--unknown", "PRIVATE_ARG"],
        ["--expected-count", "PRIVATE_ARG"],
        ["--expected-count", "0"],
        ["--labels", "PRIVATE_MISSING.xlsx"],
        ["--sheet", "PRIVATE_SHEET"],
    ],
)
def test_cli_errors_are_sanitized(dataset, args, capsys):
    assert invoke(dataset, *args)[0] == 2
    captured = capsys.readouterr()
    assert "PRIVATE_" not in captured.out + captured.err
    assert "input=" in captured.err


def test_docx_dependency_is_preloaded_before_hash(
    dataset, author_package, monkeypatch, capsys
):
    folder, _, _ = author_package
    original = audit.importlib.import_module

    def importing(name):
        if name == "docx":
            raise ModuleNotFoundError("PRIVATE_MODULE_PATH")
        return original(name)

    monkeypatch.setattr(audit.importlib, "import_module", importing)
    monkeypatch.setattr(
        audit, "sha256", lambda _: pytest.fail("hashing preceded dependency preflight")
    )
    assert invoke(dataset, "--taxonomy", folder / "taxonomy.docx")[0] == 2
    err = capsys.readouterr().err
    assert "input=taxonomy" in err and "module=docx" in err and "PRIVATE" not in err


def test_input_changed_during_audit_exit_one(dataset, monkeypatch):
    original = audit.read_labels

    def changing(path, sheet, collector):
        result = original(path, sheet, collector)
        with path.open("ab") as stream:
            stream.write(b"synthetic mutation after loading")
        return result

    monkeypatch.setattr(audit, "read_labels", changing)
    code, report = invoke(dataset)
    assert (
        code == 1
        and report["summary"]["issue_codes"]["input_changed_during_audit"] == 1
    )


def test_video_changes_between_hash_and_check(dataset, monkeypatch):
    original = audit.sha256

    def changing(path):
        result = original(path)
        if path.name == "X-1.mov":
            path.write_bytes(b"changed synthetic bytes")
        return result

    monkeypatch.setattr(audit, "sha256", changing)
    code, report = invoke(dataset)
    assert (
        code == 1 and report["summary"]["issue_codes"]["file_changed_during_hash"] == 1
    )


def test_streaming_hash_detects_mutation_while_read(dataset, monkeypatch):
    path = dataset[2] / "X-1.mov"
    path.write_bytes(b"x" * (2 * 1024 * 1024))
    original = hashlib.sha256

    class MutatingHash:
        def __init__(self):
            self.inner = original()
            self.done = False

        def update(self, chunk):
            self.inner.update(chunk)
            if not self.done:
                self.done = True
                with path.open("ab") as stream:
                    stream.write(b"mutation")

        def hexdigest(self):
            return self.inner.hexdigest()

    monkeypatch.setattr(audit.hashlib, "sha256", MutatingHash)
    with pytest.raises(audit.InputChanged, match="file_changed_during_hash"):
        audit.sha256(path)


@pytest.mark.parametrize(
    "change,code",
    [
        ("hash", "hash_mismatch"),
        ("size", "size_mismatch"),
        ("bounds", "timestamp_out_of_bounds"),
        ("label", "manifest_label_mismatch"),
        ("time", "manifest_timestamp_mismatch"),
        ("duration", "invalid_manifest_duration"),
    ],
)
def test_manifest_errors_exit_one(dataset, author_package, change, code):
    folder, _, manifest = author_package
    field, value = {
        "hash": ("sha256", "0" * 64),
        "size": ("size_bytes", 0),
        "bounds": ("duration", 8),
        "label": ("label", "без нарушения"),
        "time": ("timestamp", "not a time"),
        "duration": ("duration", -1),
    }[change]
    manifest[0][field] = value
    (folder / "manifest.json").write_text(json.dumps(manifest))
    status, report = invoke(dataset, "--manifest", folder / "manifest.json")
    assert status == 1 and report["summary"]["issue_codes"][code] >= 1
    if change == "hash":
        assert report["videos"][0]["timestamp_bounds"] == "unverified"


def test_missing_manifest_hash_and_bad_static_script_warn(dataset, author_package):
    folder, _, manifest = author_package
    manifest[0].pop("sha256")
    (folder / "manifest.json").write_text(json.dumps(manifest))
    (folder / "broken.py").write_text("this is not ( valid Python")
    code, report = invoke(dataset, "--manifest", folder / "manifest.json")
    assert code == 0
    assert {"missing_manifest_hash", "package_script_ast_error"} <= set(
        report["summary"]["issue_codes"]
    )


def test_taxonomy_only_and_report_only_supported(dataset, author_package):
    folder, _, _ = author_package
    assert invoke(dataset, "--taxonomy", folder / "taxonomy.docx")[0] == 0
    code, report = invoke(dataset, "--report-data", folder / "report.json")
    assert (
        code == 0 and report["taxonomy"]["registry_comparison"]["status"] == "skipped"
    )


@pytest.mark.parametrize("headings", ["missing", "wrong_ids"])
def test_taxonomy_only_checks_generic_ids_even_with_a_valid_registry(dataset, headings):
    document = Document()
    if headings == "wrong_ids":
        for prefix, count in (("Т", 17), ("Д", 15)):
            for number in range(2, count + 2):
                document.add_heading(f"{prefix}{number:02} Synthetic", 2)
    table = document.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Видео и метка"
    for name in ("X-1", "X-2"):
        table.add_row().cells[0].text = name
    path = dataset[0] / "taxonomy-only.docx"
    document.save(path)
    code, report = invoke(dataset, "--taxonomy", path)
    assert code == 1
    assert report["taxonomy"]["registry_comparison"]["matched"] == 2
    assert report["summary"]["issue_codes"]["id_set_mismatch"] == 2
    if headings == "missing":
        assert report["summary"]["issue_codes"]["unexpected_count"] == 2


@pytest.mark.parametrize(
    "mutation,expected",
    [
        ("sections", "invalid_report_sections"),
        ("array", "report_section_must_be_array"),
        ("width", "report_row_array_width_mismatch"),
        ("field", "report_row_field_type_mismatch"),
        ("refs", "report_source_ids_must_be_array"),
        ("registry", "report_registry_must_be_array"),
        ("registry_fields", "invalid_report_registry_schema"),
    ],
)
def test_optional_report_schema_failures(dataset, author_package, mutation, expected):
    folder, raw, _ = author_package
    altered = copy.deepcopy(raw)
    if mutation == "sections":
        altered.pop("sources")
    elif mutation == "array":
        altered["sources"] = {}
    elif mutation == "width":
        altered["types"][0] = {}
    elif mutation == "field":
        altered["types"][0][1] = 123
    elif mutation == "refs":
        altered["types"][0][-1] = [1]
    elif mutation == "registry":
        altered["registry"] = {}
    else:
        altered["registry"][0].pop("summary")
    with pytest.raises(ValueError, match=expected):
        audit.taxonomy_audit(altered, None, None, None, {}, 2, audit.Audit())


def test_schema_and_name_collisions_are_not_silent(dataset, author_package):
    folder, raw, _ = author_package
    raw["types"][0][-1] = ["И9"]
    raw["registry"][1]["id"] = 1
    raw["registry"][1]["name"] = "X-1"
    (folder / "report.json").write_text(json.dumps(raw))
    code, report = invoke(dataset, *package_args(folder))
    assert code == 1
    assert {
        "unknown_source_reference",
        "duplicate_numeric_id",
        "exact_duplicate",
    } <= set(report["summary"]["issue_codes"])


@pytest.mark.parametrize(
    "raw,expected",
    [('{"a":1,"a":2}', "duplicate_json_key"), ('{"a":NaN}', "nonfinite_json_number")],
)
def test_json_strictness(dataset, raw, expected):
    p = dataset[0] / "bad.json"
    p.write_text(raw)
    with pytest.raises(ValueError, match=expected):
        audit.read_json(p)


def test_visual_and_manifest_structure_errors(dataset):
    path = dataset[0] / "bad.json"
    for raw, code in [
        ({}, "visual_file_must_be_array"),
        ([{}], "invalid_visual_schema"),
    ]:
        path.write_text(json.dumps(raw))
        with pytest.raises(ValueError, match=code):
            audit.visual_ids([path], {}, 2, audit.Audit())
    with pytest.raises(ValueError, match="manifest_must_be_array"):
        audit.manifest_rows({}, audit.Audit())
    collector = audit.Audit()
    audit.manifest_rows(
        [
            {},
            {"file": "bad", "name": "bad"},
            {"file": "X-1.mov", "name": "X-2", "id": 7},
        ],
        collector,
    )
    assert {
        "invalid_manifest_row",
        "invalid_manifest_name",
        "manifest_name_mismatch",
        "manifest_id_mismatch",
    } <= {i["code"] for i in collector.issues}


def test_empty_invalid_docx_and_header_exit_two(dataset, author_package):
    folder, _, _ = author_package
    doc = Document()
    doc.save(folder / "empty.docx")
    assert invoke(dataset, "--taxonomy", folder / "empty.docx")[0] == 2
    book = Workbook()
    book.active.append(["wrong", "headers", "only"])
    book.save(dataset[1])
    assert invoke(dataset)[0] == 2


def test_name_and_count_errors_exit_one(dataset):
    _, _, videos = dataset
    (videos / "Х-1.mov").write_bytes(b"other synthetic fixture")
    status, report = invoke(dataset)
    assert status == 1
    assert {"normalized_name_collision", "unexpected_count", "id_set_mismatch"} <= set(
        report["summary"]["issue_codes"]
    )
